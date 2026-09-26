import pytest
from docx import Document

from app.services.document_processor import extract_text
from app.utils.exceptions import (
    UnsupportedFileFormatError,
    EmptyDocumentError,
    CorruptedDocumentError,
    DocumentProcessingError,
)


@pytest.fixture
def sample_docx(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("Skills: Python, Docker, AWS")
    doc.save(path)
    return str(path)


@pytest.fixture
def empty_docx(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    return str(path)


@pytest.fixture
def unsupported_file(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("hello")
    return str(path)


@pytest.fixture
def corrupted_pdf(tmp_path):
    path = tmp_path / "fake.pdf"
    path.write_text("this is not a real pdf")
    return str(path)


def test_extract_text_valid_docx(sample_docx):
    text = extract_text(sample_docx)
    assert "Jane Smith" in text
    assert "Docker" in text


def test_extract_text_unsupported_format(unsupported_file):
    with pytest.raises(UnsupportedFileFormatError):
        extract_text(unsupported_file)


def test_extract_text_empty_document(empty_docx):
    with pytest.raises(EmptyDocumentError):
        extract_text(empty_docx)


def test_extract_text_corrupted_pdf(corrupted_pdf):
    with pytest.raises(CorruptedDocumentError):
        extract_text(corrupted_pdf)


def test_extract_text_missing_file():
    with pytest.raises(DocumentProcessingError):
        extract_text("tests/sample_data/does_not_exist.pdf") 